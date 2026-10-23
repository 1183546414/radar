# -*- coding: utf-8 -*-
"""生成 ICDT 算法介绍 Word 文档（含网络结构图）。

用法:
    D:\\conda\\envs\\radar\\python.exe outputs/20261013_icdt/code/make_icdt_docx.py
产物:
    outputs/20261013_icdt/ICDT算法介绍.docx
    outputs/20261013_icdt/icdt_structure.png（结构图，可单独复用）
"""
from __future__ import annotations

from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch

HERE = Path(__file__).resolve()
ROUND = HERE.parents[1]

plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei"]
plt.rcParams["axes.unicode_minus"] = False

C_OBS, C_NB, C_Q, C_C = "#7FBF7F", "#6FA8DC", "#B39DDB", "#F6B26B"
C_ENC, C_HEAD, C_ENS = "#D5A6BD", "#FFE599", "#D9D9D9"


def draw_structure(path: Path):
    """ICDT 网络结构图：序列 token 化 → Transformer 编码器 → 打分头 → 悲观集成。"""
    fig, ax = plt.subplots(figsize=(13.2, 6.6), dpi=200)
    ax.set_xlim(0, 132)
    ax.set_ylim(0, 66)
    ax.axis("off")

    def box(x, y, w, h, fc, text, fs=10, ec="#555555", lw=1.0, bold=False):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.4",
                                    fc=fc, ec=ec, lw=lw))
        ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=fs,
                fontweight="bold" if bold else "normal", wrap=True)

    # ---- 第 1 层：输入序列 token 化 ----
    ax.text(2, 60.5, "① 序列构造（长度 15，四类 token，统一 143 维特征 [a; r; sim; z] + 类型嵌入）",
            fontsize=11.5, fontweight="bold")
    layout = ([("obs", C_OBS)] * 8 + [("nb", C_NB)] * 5 + [("z", C_Q), ("a", C_C)])
    x0, w_t, gap = 3.0, 5.2, 0.75
    for i, (lab, c) in enumerate(layout):
        x = x0 + i * (w_t + gap)
        box(x, 52, w_t, 5.6, c, lab, fs=9)
    ax.text(x0, 50.2, "场景内观测 (a_i, r_i)×m", fontsize=8.5, color="#38761D")
    ax.text(x0 + 8 * (w_t + gap), 50.2, "邻居最优经验 (a*,r*,sim)×5",
            fontsize=8.5, color="#1C4587")
    ax.text(x0 + 13 * (w_t + gap) - 3, 50.2, "查询 z / 候选 a", fontsize=8.5,
            color="#4B2E83")
    ax.text(x0, 47.3, "m=0..8：k-shot 观测，不足补零+mask", fontsize=8.5, color="#666666")

    # ---- 第 2 层：嵌入 + Transformer ----
    ax.text(2, 40.5, "② 嵌入与编码", fontsize=11.5, fontweight="bold")
    box(3, 30, 17, 7.5, "#EFEFEF", "线性嵌入\nLinear 143→128", fs=9.5)
    box(27, 24.5, 44, 13, C_ENC,
        "Transformer 编码器 × 4 层（Pre-Norm）\n"
        "双向自注意力（集合推理，非因果）\n"
        "d_model=128, heads=4, FF=256, dropout=0.1\n"
        "key_padding_mask 屏蔽补零槽位", fs=9.5)
    ax.add_patch(FancyArrowPatch((20.5, 33.7), (26.5, 31), arrowstyle="-|>",
                                 mutation_scale=14, color="#444444"))

    # ---- 第 3 层：预测头 + 集成 ----
    ax.text(2, 19.5, "③ 打分与悲观集成", fontsize=11.5, fontweight="bold")
    box(27, 10, 26, 7.5, C_HEAD, "取候选 token 位置输出\nMLP 128→64→1 → R(预测)", fs=9.5)
    ax.add_patch(FancyArrowPatch((49, 31), (40, 18), arrowstyle="-|>",
                                 mutation_scale=14, color="#444444"))
    box(62, 8.5, 46, 10.5, C_ENS,
        "5-seed 集成（42/1042/2042/3042/4042）\n"
        "μ(a) = 5 成员均值，σ(a) = 5 成员标准差\n"
        "悲观打分 score(a) = μ(a) − 1.0·σ(a)\n"
        "推荐 a* = argmax score(a)", fs=9.5, bold=False)
    ax.add_patch(FancyArrowPatch((53.5, 13.7), (61.5, 13.7), arrowstyle="-|>",
                                 mutation_scale=14, color="#444444"))

    # ---- 底部注释：信息流 ----
    box(3, 0.5, 105, 5.8, "#FFFFFF",
        "信息流：邻居经验/同场景观测（类比参照）×  查询状态 z（定位）×  候选动作 a（被评对象）"
        "  →  上下文条件化的相对打分", fs=9, ec="#999999")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def build_docx(png: Path, out: Path):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    def h(text, level=1):
        p = doc.add_heading(text, level=level)
        for r in p.runs:
            r.font.name = "Calibri"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
            r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)
        return p

    def para(text, bold=False, size=10.5):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
        return p

    def bullet(text):
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(text)
        r.font.size = Pt(10.5)
        return p

    # ---- 标题 ----
    t = doc.add_heading("ICDT：上下文决策 Transformer（In-Context Decision Transformer）",
                        level=0)
    for r in t.runs:
        r.font.name = "Calibri"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    p = doc.add_paragraph()
    r = p.add_run("雷达抗干扰动作推荐 · 小样本上下文适配 · Round 20261013")
    r.italic = True
    r.font.size = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ---- 1 概述 ----
    h("1. 算法概述", 1)
    para("ICDT 解决的问题是：给定当前干扰状态 S，从候选动作池（检索池 ∪ 随机采样，约 200 个）"
         "中选出使己方 SINR 最高的雷达动作 A。与既有 Retrieve+Q 管线的本质区别在于信息利用方式："
         "旧管线把检索到的相似场景仅用作“状态定位”（邻居状态向量加权均值作为条件），邻居场景中"
         "已被验证的最优动作从未进入推理；ICDT 将两类证据统一构造为 token 序列——跨场景的邻居"
         "经验（最优动作、得分、相似度）与场景内的少量实测反馈 (aᵢ, rᵢ)——交给 Transformer 做"
         "双向自注意力，在前向传播中完成“经验类比 + 局部参照系”下的候选打分，全程零梯度适配。")
    para("核心特性：", bold=True)
    bullet("经验检索：邻居场景的“标准答案”（最优动作与得分）显式进入上下文，模型学类比而非重新学映射")
    bullet("相对打分：候选动作在“同场景已知观测”构成的局部参照系中被评估，抵消场景级偏置")
    bullet("自适应融合：注意力权重自动平衡邻居经验、实测反馈与状态条件——检索证据弱时自动降权邻居")
    bullet("k-shot 内建：训练时观测数 m∈0..8 随机，推理时观测可从 0 递增到 8，预算-性能连续可调")

    # ---- 2 网络结构 ----
    h("2. 网络结构", 1)
    doc.add_picture(str(png), width=Cm(16.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para("图 1  ICDT 网络结构与数据流（序列构造 → Transformer 编码器 → 打分头 → 5-seed 悲观集成）",
         size=9).alignment = WD_ALIGN_PARAGRAPH.CENTER

    h("2.1 序列构造与 token 化", 2)
    para("每条输入序列固定长度 15，全部 token 统一映射为 143 维特征："
         "[a(9); r(1); sim(1); z(128)]（缺失分量补零）+ 4 维类型 one-hot，无位置嵌入"
         "（集合无序，token 身份由类型嵌入区分）。")
    tbl = doc.add_table(rows=5, cols=4)
    tbl.style = "Light Grid Accent 1"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    heads = ["Token 类型", "数量", "内容", "语义角色"]
    rows = [
        ("obs 观测", "0~8 槽", "(aᵢ, rᵢ) 当前场景已评估动作及真值", "k-shot 实测反馈，局部参照系"),
        ("nb 邻居经验", "5 槽", "同类 Top-5 邻居场景的最优动作 (a*, r*) 与相似度 sim", "跨场景经验类比（经验检索）"),
        ("query 查询", "1 槽", "当前干扰状态潜向量 z ∈ R¹²⁸", "场景定位条件"),
        ("cand 候选", "1 槽", "待评候选动作 a ∈ R⁹", "被评对象，其输出位置即预测位"),
    ]
    for j, s in enumerate(heads):
        tbl.rows[0].cells[j].text = s
    for i, row in enumerate(rows, start=1):
        for j, s in enumerate(row):
            tbl.rows[i].cells[j].text = s

    h("2.2 Transformer 编码器", 2)
    para("4 层 Pre-Norm TransformerEncoderLayer（d_model=128，多头注意力 heads=4，"
         "前馈 FF=256，dropout=0.1），采用双向自注意力而非因果掩码——任务是对固定候选集合做"
         "联合推理而非自回归生成：候选 token 可聚合全部上下文信息，观测/邻居/查询 token 之间"
         "也互相通信。变长观测通过 key_padding_mask 屏蔽补零槽位。模型仅约 0.56M 参数。")

    h("2.3 打分头与悲观集成", 2)
    para("取候选 token 位置的输出，经 MLP(128→64→1) 输出标准化 R 预测 R̂。训练 5 个不同随机"
         "种子的成员，推理时对候选池打分取 μ(a) 与 σ(a)，悲观打分 score(a)=μ(a)−1.0·σ(a)，"
         "argmax 选动作——成员分歧大的候选被自动压低，抑制过自信。")

    # ---- 3 训练 ----
    h("3. 训练方式", 1)
    bullet("数据：Round2 信号级数据（5000 场景 × 每场景 20 动作 = 10 万行，干扰状态 z 由冻结 VAE 编码）")
    bullet("k-shot 模拟：每样本随机取 1 行做查询（候选+标签），其余 19 行中无放回采 m∈{0..8} 行"
           "做观测 token；邻居经验 token 来自同类 Top-5（余弦相似度）场景的 argmax-R 动作")
    bullet("损失：MSE（标准化 R 空间）；Adam lr=3e-4，batch=256，200 步/epoch")
    bullet("按场景 80/20 划分（与 Round2 同协议，防同场景泄漏）；固定验证集 4096 样本，"
           "PATIENCE=25 早停（5 个成员分别于 42~56 epoch 停止，val MSE 0.180~0.185，R²≈0.82）")
    bullet("向量化构造：随机键排序实现观测的无放回采样（查询行必不在观测中），每步动态重采样")

    # ---- 4 推理 ----
    h("4. 推理与 k-shot 适配循环", 1)
    para("对每个新干扰场景：构造候选池（检索池 ∪ 随机 100A）→ 检索同类 Top-5 邻居经验 → "
         "按可用预算 m 获取观测（真值来自信号级仿真/实测反馈）→ 构造序列打分 → 悲观 argmax。"
         "三种预算模式：")
    bullet("零观测（m=0）：纯上下文（邻居经验 + z 条件），一次前向，零真值开销")
    bullet("随机选点：候选池随机抽 m 个动作获取真值，滚入上下文后统一打分")
    bullet("主动选点：迭代“打分 → 选未仿真候选中悲观分最高者 → 仿真 → 滚入上下文”共 m 轮，"
           "适配预算花在模型分歧最大的位置")

    # ---- 5 实验结果 ----
    h("5. 实验结果摘要", 1)
    para("协议：50 个留出场景、seed=2026、信号级真值（JSR=10dB），与既有各轮完全同条件。"
         "指标为相对 Oracle-K20 的 regret（dB，负值=超过该上界）。")
    tbl2 = doc.add_table(rows=6, cols=4)
    tbl2.style = "Light Grid Accent 1"
    data2 = [
        ("方法", "regret (dB)", "零 regret", "胜率 vs Q-only"),
        ("Retrieve+Q-Union（Round2 基线）", "−5.08", "84%", "90%"),
        ("Q-Ens-悲观（前章默认）", "−4.97", "82%", "96%"),
        ("ICDT 零观测（m=0）", "−5.08", "90%", "96%"),
        ("ICDT 随机观测 m=5", "−5.21（全场最佳）", "88%", "94%"),
        ("Oracle-Wide（100 随机 A 上界）", "−5.73", "92%", "96%"),
    ]
    for i, row in enumerate(data2):
        for j, s in enumerate(row):
            tbl2.rows[i].cells[j].text = s
    para("")
    bullet("零观测即追平最优基线均值且零 regret 最高——增益来自上下文经验表达本身，不依赖真值开销")
    bullet("m=5 时 −5.21 dB 超过基线与前章最优 0.13~0.24 dB，每场景仅 5 次真值查询")
    bullet("检索证据最弱的 hard 桶（9 场景）：主动选点 m=3 达 −2.61 dB，超过前章最优 −2.54、"
           "基线 −1.68，逼近 Oracle-Wide 的 −2.71")
    bullet("5 成员 val MSE 极稳（std≈0.002），悲观集成保持 96% 配对胜率")

    # ---- 6 对比 ----
    h("6. 与既有方法的关系", 1)
    para("ICDT 不替代检索与前章 Q-Ensemble，而是把它们变成自身组件：检索库提供邻居经验 token，"
         "悲观集成思想（μ−λσ）原样复用。与前章 QRetrievalNet 的差别是范式级的：旧模型做全局"
         "绝对回归（邻居仅定位状态），ICDT 做上下文条件化的相对打分（邻居的答案直接可见）。"
         "主动选点循环（BO 式预算分配）也被统一进同一个 Transformer 叙事：由模型悲观分驱动的"
         "主动上下文适配。")

    doc.save(out)


if __name__ == "__main__":
    png = ROUND / "icdt_structure.png"
    draw_structure(png)
    print(f"结构图: {png}")
    out = ROUND / "ICDT算法介绍.docx"
    build_docx(png, out)
    print(f"文档: {out}")
